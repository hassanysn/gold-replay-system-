"""Build Monthly_Budget_Tracker.xlsx — a fill-in-the-yellow-cells monthly budget workbook.

Usage:  python build_budget_tracker.py [output_path]

Every number on Dashboard / Daily Flow is a live Excel formula driven by the
Setup and Transactions sheets. The script writes the workbook with openpyxl,
then (if LibreOffice is available) recalculates a scratch copy and copies the
computed values back in as cached results, so previewers that don't calculate
still show numbers. Excel recalculates everything on open regardless.
"""
import datetime as dt
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines, DateAxis
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint, SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

OUT = sys.argv[1] if len(sys.argv) > 1 else "Monthly_Budget_Tracker.xlsx"

# ---------------------------------------------------------------- styling
FONT = "Arial"
ACCENT = "1C5CAB"        # headers, titles
ACCENT_MID = "2A78D6"    # chart marks, data bars
ACCENT_SOFT = "EAF1FB"   # KPI card background
INPUT_FILL = "FFF4C7"    # light yellow  = type here
AUTO_FILL = "EDEDED"     # grey          = automatic
INK = "222222"
MUTED = "6B6B6B"
WHITE = "FFFFFF"
# categorical slots for the doughnut (validated CVD-safe order), "Other" in grey
DONUT_COLORS = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "A3A3A3"]

MONEY = "#,##0.00"
MONEY_RED = '#,##0.00;[Red]-#,##0.00'
DATE = "dd-mmm-yyyy"

fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="D0D0D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
UNLOCKED = Protection(locked=False)


def put(ws, ref, value=None, *, size=10, bold=False, italic=False, color=INK, bg=None,
        fmt=None, h=None, v="center", wrap=False, border=None, locked=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if border:
        c.border = border
    if locked is False:
        c.protection = UNLOCKED
    return c


def money(x):
    """Text like $1,234.56 / -$12.00 using the currency symbol from Setup (locale-safe FIXED)."""
    return f'IF({x}<0,"-","")&CurSym&FIXED(ABS({x}),2)'


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)

ws_start = wb.active
ws_start.title = "Start Here"
ws_setup = wb.create_sheet("Setup")
ws_tx = wb.create_sheet("Transactions")
ws_dash = wb.create_sheet("Dashboard")
ws_flow = wb.create_sheet("Daily Flow")

ws_start.sheet_properties.tabColor = ACCENT
ws_setup.sheet_properties.tabColor = "F2C200"
ws_tx.sheet_properties.tabColor = "F2C200"
ws_dash.sheet_properties.tabColor = ACCENT
ws_flow.sheet_properties.tabColor = "A6A6A6"

# ---------------------------------------------------------------- named ranges
names = {
    "StartDate": "Setup!$C$5",
    "EndDate": "Setup!$C$6",
    "DaysInMonth": "Setup!$C$7",
    "CurSym": "Setup!$C$9",
    "ExpectedIncome": "Setup!$J$5",
    "TotalBudget": "Setup!$J$6",
    # structured references grow with the tables, so dropdowns pick up new rows
    "IncomeSources": "tblIncome[Source]",
    "ExpenseCategories": "tblCategories[Category]",
    "PaymentMethods": "tblPayments[Method]",
}
for n, ref in names.items():
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def make_table(ws, name, ref, columns, calculated=None):
    t = Table(displayName=name, ref=ref)
    t.tableColumns = []
    for i, col in enumerate(columns, 1):
        tc = TableColumn(id=i, name=col)
        if calculated and col in calculated:
            tc.calculatedColumnFormula = TableFormula(attr_text=calculated[col])
        t.tableColumns.append(tc)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False,
                                      showColumnStripes=False, showFirstColumn=False,
                                      showLastColumn=False)
    ws.add_table(t)
    return t


def dv(ws, rng, **kw):
    d = DataValidation(allow_blank=True, showErrorMessage=True, showInputMessage=True, **kw)
    ws.add_data_validation(d)
    for part in rng.split():
        d.add(part)
    return d


# ================================================================ SETUP
ws = ws_setup
ws.sheet_view.showGridLines = False
widths(ws, {"A": 2, "B": 26, "C": 16, "D": 3, "E": 26, "F": 16, "G": 12, "H": 3, "I": 26, "J": 14})
put(ws, "B1", "Setup", size=18, bold=True, color=ACCENT)
put(ws, "B2", "Fill in the yellow cells. Grey cells calculate on their own. "
              "To add a row to any table, just type in the empty row directly below it.",
    color=MUTED, italic=True)

put(ws, "B4", "1. Month & currency", size=12, bold=True, color=ACCENT)
rows = [
    (5, "Month start date", dt.date(2026, 11, 1), DATE, True, "Type the 1st of the month, e.g. 1-Nov-2026"),
    (6, "Month end date", "=EOMONTH(C5,0)", DATE, False, "Automatic (last day of that month)"),
    (7, "Days in month", "=C6-C5+1", "0", False, "Automatic"),
    (8, "Month being tracked", "=C5", "mmmm yyyy", False, "Automatic"),
    (9, "Currency symbol", "$", "@", True, "Any symbol or code: $, €, £, ₹, AED …"),
]
for r, label, val, fmt, is_input, note in rows:
    put(ws, f"B{r}", label, bold=True, border=BOX)
    put(ws, f"C{r}", val, fmt=fmt, h="center", border=BOX,
        bg=INPUT_FILL if is_input else AUTO_FILL, locked=False if is_input else None,
        bold=is_input)
    put(ws, f"E{r}", note, size=9, italic=True, color=MUTED)
dv(ws, "C5", type="custom", formula1="AND(ISNUMBER(C5),DAY(C5)=1)",
   errorTitle="Use the 1st of the month",
   error="Please enter the FIRST day of the month you want to track, e.g. 1-Nov-2026.",
   promptTitle="Month start date", prompt="Enter the 1st day of the month (e.g. 1-Nov-2026).")

put(ws, "I4", "At a glance (automatic)", size=12, bold=True, color=ACCENT)
for r, label, f in [
    (5, '="Expected income ("&CurSym&")"', "=SUM(tblIncome[Expected Amount])"),
    (6, '="Total monthly budget ("&CurSym&")"', "=SUM(tblCategories[Monthly Budget])"),
    (7, '="Planned surplus ("&CurSym&")"', "=J5-J6"),
]:
    put(ws, f"I{r}", label, bold=True, border=BOX)
    put(ws, f"J{r}", f, fmt=MONEY_RED, h="right", bg=AUTO_FILL, border=BOX)

# --- tables (side by side so each can grow downwards freely)
put(ws, "B11", "2. Income sources", size=12, bold=True, color=ACCENT)
put(ws, "B12", "Where your money comes from and how much you expect.", size=9, italic=True, color=MUTED)
put(ws, "E11", "3. Expense categories", size=12, bold=True, color=ACCENT)
put(ws, "E12", "Budget per category. The first 20 show on the Dashboard.", size=9, italic=True, color=MUTED)
put(ws, "I11", "4. Payment methods", size=12, bold=True, color=ACCENT)
put(ws, "I12", "How you pay.", size=9, italic=True, color=MUTED)

income = [("Salary", 4500), ("Freelance", 800), ("Other", 0)]
categories = [
    ("Rent", 1400, "Fixed"), ("Utilities", 180, "Variable"), ("Internet & Phone", 90, "Fixed"),
    ("Groceries", 500, "Variable"), ("Transport/Fuel", 220, "Variable"), ("Dining Out", 200, "Variable"),
    ("Health", 100, "Variable"), ("Shopping", 200, "Variable"), ("Entertainment", 120, "Variable"),
    ("Subscriptions", 45, "Fixed"), ("Education", 100, "Fixed"), ("Savings", 600, "Fixed"),
    ("Debt Payments", 350, "Fixed"), ("Gifts/Family", 100, "Variable"), ("Miscellaneous", 100, "Variable"),
]
methods = ["Cash", "Card", "Bank Transfer", "Wallet"]

for col, head in zip("BC", ["Source", "Expected Amount"]):
    put(ws, f"{col}13", head, bold=True, color=WHITE, bg=ACCENT)
for i, (src, amt) in enumerate(income, 14):
    put(ws, f"B{i}", src, bg=INPUT_FILL, locked=False)
    put(ws, f"C{i}", amt, fmt=MONEY, bg=INPUT_FILL, locked=False)
make_table(ws, "tblIncome", f"B13:C{13 + len(income)}", ["Source", "Expected Amount"])

for col, head in zip("EFG", ["Category", "Monthly Budget", "Type"]):
    put(ws, f"{col}13", head, bold=True, color=WHITE, bg=ACCENT)
for i, (cat, bud, typ) in enumerate(categories, 14):
    put(ws, f"E{i}", cat, bg=INPUT_FILL, locked=False)
    put(ws, f"F{i}", bud, fmt=MONEY, bg=INPUT_FILL, locked=False)
    put(ws, f"G{i}", typ, h="center", bg=INPUT_FILL, locked=False)
make_table(ws, "tblCategories", f"E13:G{13 + len(categories)}", ["Category", "Monthly Budget", "Type"])

put(ws, "I13", "Method", bold=True, color=WHITE, bg=ACCENT)
for i, m in enumerate(methods, 14):
    put(ws, f"I{i}", m, bg=INPUT_FILL, locked=False)
make_table(ws, "tblPayments", f"I13:I{13 + len(methods)}", ["Method"])

dv(ws, "C14:C300 F14:F300", type="decimal", operator="greaterThanOrEqual", formula1="0",
   errorTitle="Amount", error="Please enter a number of 0 or more.")
dv(ws, "G14:G300", type="list", formula1='"Fixed,Variable"',
   errorTitle="Type", error="Choose Fixed or Variable from the list.")
ws.freeze_panes = None

# ================================================================ TRANSACTIONS
ws = ws_tx
TX_COLS = ["Date", "Type", "Category", "Description", "Amount", "Payment Method", "Notes",
           "Day", "Week", "Signed Amount", "Rank Key"]
widths(ws, {"A": 14, "B": 11, "C": 20, "D": 34, "E": 13, "F": 17, "G": 30,
            "H": 7, "I": 8, "J": 16, "K": 10})
TX_FIRST, TX_ROWS = 5, 200
TX_LAST = TX_FIRST + TX_ROWS - 1
DV_LAST = 2000  # validation reaches past the table so it keeps working as the table grows

put(ws, "A1", "Transactions", size=18, bold=True, color=ACCENT)
put(ws, "A2", "Month being tracked:", bold=True)
put(ws, "C2", "=StartDate", fmt="mmmm yyyy", bold=True, bg=AUTO_FILL, h="center")
put(ws, "D2", "  One row per income or expense. Pick Type first, then Category.",
    size=9, italic=True, color=MUTED)
ws.merge_cells("A3:G3")
put(ws, "A3", "SAMPLE DATA  ▸  The 20 example rows below are only there to show how it works. "
              "Before you start: select cells A5:G24 and press Delete (clear the yellow columns only).",
    bold=True, color="9C0006", bg="FFE1E1", wrap=True)
ws.row_dimensions[3].height = 30

for i, head in enumerate(TX_COLS):
    put(ws, f"{chr(65 + i)}4", head, bold=True, color=WHITE, bg=ACCENT if i < 7 else "8C8C8C", h="center", wrap=True)


def this(col):
    return f"tblTx[[#This Row],[{col}]]"


calc = {
    "Day": f'IF(ISNUMBER({this("Date")}),DAY({this("Date")}),"")',
    "Week": f'IF(ISNUMBER({this("Date")}),INT((DAY({this("Date")})-1)/7)+1,"")',
    "Signed Amount": (f'IF(ISNUMBER({this("Amount")}),IF({this("Type")}="Income",{this("Amount")},'
                      f'IF({this("Type")}="Expense",-{this("Amount")},0)),"")'),
    # this month's expenses only, with a tiny row-based tie-breaker so Top-5 never repeats a row
    "Rank Key": (f'IF(AND({this("Type")}="Expense",ISNUMBER({this("Amount")}),ISNUMBER({this("Date")})),'
                 f'IF(AND({this("Date")}>=StartDate,{this("Date")}<=EndDate),'
                 f'{this("Amount")}-ROW()/10^7,""),"")'),
}

d = lambda day: dt.date(2026, 11, day)
samples = [
    (d(1), "Income", "Salary", "Monthly salary", 4500.00, "Bank Transfer", ""),
    (d(1), "Expense", "Rent", "November rent", 1400.00, "Bank Transfer", "Paid on time"),
    (d(2), "Expense", "Groceries", "Weekly groceries - supermarket", 112.40, "Card", ""),
    (d(3), "Expense", "Subscriptions", "Streaming + music", 24.98, "Card", "Auto-renew"),
    (d(4), "Expense", "Transport/Fuel", "Fuel top-up", 55.00, "Card", ""),
    (d(5), "Expense", "Utilities", "Electricity bill", 96.30, "Bank Transfer", ""),
    (d(6), "Expense", "Dining Out", "Lunch with colleagues", 28.50, "Card", ""),
    (d(7), "Expense", "Internet & Phone", "Mobile + home internet", 89.99, "Bank Transfer", ""),
    (d(8), "Expense", "Savings", "Transfer to savings account", 600.00, "Bank Transfer", "Emergency fund"),
    (d(9), "Expense", "Groceries", "Weekly groceries", 98.75, "Card", ""),
    (d(10), "Income", "Freelance", "Website project - client A", 650.00, "Bank Transfer", "Invoice #014"),
    (d(11), "Expense", "Health", "Pharmacy", 32.60, "Cash", ""),
    (d(12), "Expense", "Entertainment", "Cinema tickets", 30.00, "Wallet", ""),
    (d(14), "Expense", "Debt Payments", "Credit card payment", 350.00, "Bank Transfer", ""),
    (d(15), "Expense", "Shopping", "Winter jacket", 145.00, "Card", ""),
    (d(16), "Expense", "Groceries", "Weekly groceries", 121.30, "Card", ""),
    (d(18), "Expense", "Transport/Fuel", "Monthly bus pass", 45.00, "Wallet", ""),
    (d(20), "Expense", "Gifts/Family", "Birthday gift", 60.00, "Card", ""),
    (d(22), "Expense", "Dining Out", "Dinner out", 64.20, "Card", ""),
    (d(25), "Expense", "Shopping", "Running shoes", 89.99, "Card", "Over budget example"),
]
for r in range(TX_FIRST, TX_LAST + 1):
    row = samples[r - TX_FIRST] if r - TX_FIRST < len(samples) else (None,) * 7
    fmts = [DATE, None, None, None, MONEY, None, None]
    aligns = ["center", "center", None, None, None, None, None]
    for ci, val in enumerate(row):
        put(ws, f"{chr(65 + ci)}{r}", val, fmt=fmts[ci], h=aligns[ci], bg=INPUT_FILL, locked=False)
    put(ws, f"H{r}", "=" + calc["Day"], bg=AUTO_FILL, h="center", color=MUTED)
    put(ws, f"I{r}", "=" + calc["Week"], bg=AUTO_FILL, h="center", color=MUTED)
    put(ws, f"J{r}", "=" + calc["Signed Amount"], bg=AUTO_FILL, fmt=MONEY_RED, color=MUTED)
    put(ws, f"K{r}", "=" + calc["Rank Key"], bg=AUTO_FILL, fmt="0.00", color=MUTED)
make_table(ws, "tblTx", f"A4:K{TX_LAST}", TX_COLS, calculated=calc)
ws.column_dimensions["K"].hidden = True  # internal helper for the Top-5 list
ws.freeze_panes = "A5"
ws.row_dimensions[4].height = 30

dv(ws, f"A{TX_FIRST}:A{DV_LAST}", type="date", operator="between",
   formula1="StartDate", formula2="EndDate",
   errorTitle="Date outside this month",
   error="This date is not in the month you are tracking. Use a date between the month start "
         "and end shown on the Setup sheet (or change the month on Setup).",
   promptTitle="Date", prompt="A date inside the month set on the Setup sheet.")
dv(ws, f"B{TX_FIRST}:B{DV_LAST}", type="list", formula1='"Income,Expense"',
   errorTitle="Type", error="Choose Income or Expense from the list.",
   promptTitle="Type", prompt="Income or Expense. Choose this first.")
dv(ws, f"C{TX_FIRST}:C{DV_LAST}", type="list",
   formula1=f'IF($B{TX_FIRST}="Income",IncomeSources,ExpenseCategories)',
   errorTitle="Category", error="Pick a category from the list. Add new ones on the Setup sheet.",
   promptTitle="Category", prompt="Income → your income sources. Expense → your expense categories.")
dv(ws, f"E{TX_FIRST}:E{DV_LAST}", type="decimal", operator="greaterThan", formula1="0",
   errorTitle="Amount", error="Amount must be a positive number (no minus sign). "
                              "Type decides whether it is money in or out.",
   promptTitle="Amount", prompt="Positive number. No minus sign needed.")
dv(ws, f"F{TX_FIRST}:F{DV_LAST}", type="list", formula1="PaymentMethods",
   errorTitle="Payment method", error="Pick a method from the list. Add new ones on the Setup sheet.")

# ================================================================ DAILY FLOW
ws = ws_flow
ws.sheet_view.showGridLines = False
widths(ws, {"A": 6, "B": 15, "C": 14, "D": 14, "E": 14, "F": 18, "G": 3, "H": 13, "I": 15, "J": 14})
put(ws, "A1", "Daily Flow", size=18, bold=True, color=ACCENT)
put(ws, "A2", "Automatic: one row per day of the month, built from the Transactions sheet. Nothing to type here.",
    italic=True, color=MUTED)
put(ws, "H2", "Chart feed (automatic)", bold=True, color=MUTED)
put(ws, "H3", "Repeats the last day in short months so charts never drop to zero.",
    size=8, italic=True, color=MUTED)
heads = ["Day", "Date", '="Income ("&CurSym&")"', '="Expenses ("&CurSym&")"', '="Net ("&CurSym&")"',
         '="Running Balance ("&CurSym&")"']
for col, hd in zip("ABCDEF", heads):
    put(ws, f"{col}4", hd, bold=True, color=WHITE, bg=ACCENT, h="center", wrap=True, border=BOX)
for col, hd in zip("HIJ", ["Chart date", "Running balance", "Daily spending"]):
    put(ws, f"{col}4", hd, bold=True, color=WHITE, bg="8C8C8C", h="center", wrap=True, border=BOX)
ws.row_dimensions[4].height = 30
F0, F1 = 5, 35
for n in range(1, 32):
    r = F0 + n - 1
    put(ws, f"A{r}", n, h="center", color=MUTED, border=BOX)
    put(ws, f"B{r}", f'=IF(StartDate+A{r}-1<=EndDate,StartDate+A{r}-1,"")', fmt="ddd dd-mmm",
        h="center", border=BOX)
    put(ws, f"C{r}", f'=IF(B{r}="","",SUMIFS(tblTx[Amount],tblTx[Type],"Income",tblTx[Date],B{r}))',
        fmt=MONEY, border=BOX)
    put(ws, f"D{r}", f'=IF(B{r}="","",SUMIFS(tblTx[Amount],tblTx[Type],"Expense",tblTx[Date],B{r}))',
        fmt=MONEY, border=BOX)
    put(ws, f"E{r}", f'=IF(B{r}="","",C{r}-D{r})', fmt=MONEY_RED, border=BOX)
    prev = f"F{r - 1}+" if n > 1 else ""
    put(ws, f"F{r}", f'=IF(B{r}="","",{prev}E{r})', fmt=MONEY_RED, bold=True, border=BOX)
    put(ws, f"H{r}", f"=MIN(StartDate+A{r}-1,EndDate)", fmt="dd-mmm", h="center", color=MUTED,
        bg=AUTO_FILL, border=BOX)
    put(ws, f"I{r}", f"=INDEX($F${F0}:$F${F1},MIN(A{r},DaysInMonth))", fmt=MONEY_RED, color=MUTED,
        bg=AUTO_FILL, border=BOX)
    put(ws, f"J{r}", f"=INDEX($D${F0}:$D${F1},MIN(A{r},DaysInMonth))", fmt=MONEY, color=MUTED,
        bg=AUTO_FILL, border=BOX)
put(ws, "A36", "Total", bold=True, border=BOX, bg=ACCENT_SOFT)
put(ws, "B36", None, border=BOX, bg=ACCENT_SOFT)
for col in "CDE":
    put(ws, f"{col}36", f"=SUM({col}{F0}:{col}{F1})", fmt=MONEY_RED, bold=True, border=BOX, bg=ACCENT_SOFT)
put(ws, "F36", f"=INDEX(F{F0}:F{F1},DaysInMonth)", fmt=MONEY_RED, bold=True, border=BOX, bg=ACCENT_SOFT)
ws.freeze_panes = "A5"
ws.protection.sheet = True

# ================================================================ DASHBOARD
ws = ws_dash
ws.sheet_view.showGridLines = False
widths(ws, {"A": 2, "B": 24, "C": 12, "D": 16, "E": 13, "F": 13, "G": 10, "H": 16, "I": 3,
            "J": 6, "K": 20, "L": 13, "M": 10, "N": 12, "O": 12, "P": 10})
ws.column_dimensions["P"].hidden = True  # internal sort keys

put(ws, "B2", "Monthly Budget Dashboard", size=20, bold=True, color=ACCENT)
put(ws, "B3", "=StartDate", fmt="mmmm yyyy", size=12, bold=True, color=INK, h="left")
ws.merge_cells("C3:D3")
put(ws, "C3", "=StartDate", fmt='"from "d mmm yyyy', color=MUTED, h="left")
ws.merge_cells("E3:F3")
put(ws, "E3", "=EndDate", fmt='"to "d mmm yyyy', color=MUTED, h="left")
put(ws, "K3", "Updates automatically. Nothing to type on this sheet.", size=9, italic=True, color=MUTED)
ws.row_dimensions[2].height = 30

# --- KPI cards: (label, first col, last col, numeric formula, display formula, caption)
TOT_INC = ('SUMIFS(tblTx[Amount],tblTx[Type],"Income",'
           'tblTx[Date],">="&StartDate,tblTx[Date],"<="&EndDate)')
TOT_EXP = TOT_INC.replace('"Income"', '"Expense"')
cards = [
    ("TOTAL INCOME", "B", "B", "=" + TOT_INC, "=" + money("B8"), '="Expected: "&' + money("ExpectedIncome")),
    ("TOTAL EXPENSES", "C", "D", "=" + TOT_EXP, "=" + money("C8"), '="Budget: "&' + money("TotalBudget")),
    ("NET BALANCE", "E", "F", "=B8-C8", "=" + money("E8"),
     '=IF(E8<0,"You spent more than you earned","Income minus expenses")'),
    ("% OF INCOME SPENT", "G", "H", '=IF(B8>0,C8/B8,"")', '=IF(B8>0,FIXED(G8*100,1)&"%","—")',
     '=IF(B8=0,"No income logged yet",IF(G8>1,"Spending exceeds income","of this month\'s income"))'),
    ("DAYS LEFT IN MONTH", "J", "K", "=IF(TODAY()<StartDate,DaysInMonth,MAX(0,EndDate-TODAY()))",
     "=J8", '="of "&DaysInMonth&" days"'),
    ("AVG DAILY SPEND", "L", "M", "=IFERROR(C8/M8,0)", "=" + money("L8"),
     '="Over "&M8&IF(M8=1," day"," days")'),
    ("PROJECTED MONTH-END", "N", "O", "=B8-L8*DaysInMonth", "=" + money("N8"),
     '="If spending continues at this pace"'),
]
gap = Side(style="thick", color=WHITE)
for label, c1, c2, num, disp, cap in cards:
    for r, val, kw in [(5, label, dict(size=8, bold=True, color=MUTED)),
                       (6, disp, dict(size=17, bold=True, color=INK)),
                       (7, cap, dict(size=8, italic=True, color=MUTED))]:
        if c1 != c2:
            ws.merge_cells(f"{c1}{r}:{c2}{r}")
        put(ws, f"{c1}{r}", val, h="center", bg=ACCENT_SOFT, fmt="0" if val == "=J8" else None, **kw)
        cols = [c1] if c1 == c2 else [c1, c2]
        for col in cols:
            ws[f"{col}{r}"].fill = fill(ACCENT_SOFT)
            ws[f"{col}{r}"].border = Border(left=gap if col == c1 else None,
                                            right=gap if col == c2 else None)
        for col in cols:
            ws[f"{col}4"].fill = fill(ACCENT)
            ws[f"{col}4"].border = Border(left=gap if col == c1 else None,
                                          right=gap if col == c2 else None)
    put(ws, f"{c1}8", num, fmt="0.00")
put(ws, "M8", "=IF(OR(TODAY()<StartDate,TODAY()>EndDate),DaysInMonth,TODAY()-StartDate+1)", fmt="0")
put(ws, "A8", "calc")
ws.row_dimensions[4].height = 5
ws.row_dimensions[5].height = 20
ws.row_dimensions[6].height = 34
ws.row_dimensions[7].height = 18
ws.row_dimensions[8].hidden = True  # numeric values behind the cards

ws.conditional_formatting.add("E6", FormulaRule(formula=["$E$8<0"], font=Font(color="C00000", bold=True)))
ws.conditional_formatting.add("E6", FormulaRule(formula=["$E$8>=0"], font=Font(color="00843D", bold=True)))
ws.conditional_formatting.add("N6", FormulaRule(formula=["$N$8<0"], font=Font(color="C00000", bold=True)))
ws.conditional_formatting.add("G6", FormulaRule(formula=["AND(ISNUMBER($G$8),$G$8>1)"],
                                                font=Font(color="C00000", bold=True)))
ws.conditional_formatting.add("G6", FormulaRule(formula=["AND(ISNUMBER($G$8),$G$8>=0.8)"],
                                                font=Font(color="B26B00", bold=True)))

# --- Budget vs actual (one row per Setup category, 20 slots)
put(ws, "B10", "Budget vs Actual by category", size=13, bold=True, color=ACCENT)
bv_heads = ["Category", "Type", '="Budget ("&CurSym&")"', '="Spent ("&CurSym&")"',
            '="Remaining ("&CurSym&")"', "% Used", "Progress"]
for col, hd in zip("BCDEFGH", bv_heads):
    put(ws, f"{col}11", hd, bold=True, color=WHITE, bg=ACCENT, h="center", border=BOX)
BV0, NCAT = 12, 20
BV1 = BV0 + NCAT - 1
for i in range(1, NCAT + 1):
    r = BV0 + i - 1
    put(ws, f"B{r}", f'=IFERROR(IF(INDEX(tblCategories[Category],{i})="","",'
                     f'INDEX(tblCategories[Category],{i})),"")', border=BOX)
    put(ws, f"C{r}", f'=IF(B{r}="","",IFERROR(INDEX(tblCategories[Type],{i})&"",""))',
        h="center", color=MUTED, border=BOX)
    put(ws, f"D{r}", f'=IF(B{r}="","",N(INDEX(tblCategories[Monthly Budget],{i})))', fmt=MONEY, border=BOX)
    put(ws, f"E{r}", f'=IF(B{r}="","",SUMIFS(tblTx[Amount],tblTx[Type],"Expense",tblTx[Category],B{r},'
                     f'tblTx[Date],">="&StartDate,tblTx[Date],"<="&EndDate))', fmt=MONEY, border=BOX)
    put(ws, f"F{r}", f'=IF(B{r}="","",D{r}-E{r})', fmt=MONEY_RED, border=BOX)
    put(ws, f"G{r}", f'=IF(B{r}="","",IF(D{r}>0,E{r}/D{r},IF(E{r}>0,"No budget",0)))',
        fmt="0%", h="center", border=BOX)
    put(ws, f"H{r}", f'=IF(ISNUMBER(G{r}),MIN(G{r},1),IF(G{r}="No budget",1,""))', fmt=";;;", border=BOX)
    # sort key: spent (+ tiny tie-breaker keeping Setup order); empty slots sink to the bottom
    put(ws, f"P{r}", f'=IF(B{r}="",-{i},E{r}+({NCAT + 1}-{i})/10^6)', fmt="0.000000")
TR = BV1 + 1  # totals row
put(ws, f"B{TR}", "Total", bold=True, border=BOX, bg=ACCENT_SOFT)
put(ws, f"C{TR}", None, border=BOX, bg=ACCENT_SOFT)
put(ws, f"D{TR}", f"=SUM(D{BV0}:D{BV1})", fmt=MONEY, bold=True, border=BOX, bg=ACCENT_SOFT)
put(ws, f"E{TR}", f"=SUM(E{BV0}:E{BV1})", fmt=MONEY, bold=True, border=BOX, bg=ACCENT_SOFT)
put(ws, f"F{TR}", f"=D{TR}-E{TR}", fmt=MONEY_RED, bold=True, border=BOX, bg=ACCENT_SOFT)
put(ws, f"G{TR}", f'=IF(D{TR}>0,E{TR}/D{TR},"")', fmt="0%", h="center", bold=True, border=BOX, bg=ACCENT_SOFT)
put(ws, f"H{TR}", f'=IF(ISNUMBER(G{TR}),MIN(G{TR},1),"")', fmt=";;;", border=BOX, bg=ACCENT_SOFT)
put(ws, f"B{TR + 1}", f'=IF(ABS(C8-E{TR})>0.005,"Note: "&{money(f"C8-E{TR}")}&" of this month\'s expenses '
                      f'use a category that is not in the first 20 Setup categories. Check the Category '
                      f'column on Transactions.","")', size=9, italic=True, color="C00000")

pct = f"$G{BV0}"
rng = f"G{BV0}:G{TR}"
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($B{BV0}<>"",OR({pct}="No budget",AND(ISNUMBER({pct}),{pct}>1)))'],
    fill=fill("FFC7CE"), font=Font(color="9C0006", bold=True), stopIfTrue=True))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($B{BV0}<>"",ISNUMBER({pct}),{pct}>=0.8)'],
    fill=fill("FFEB9C"), font=Font(color="9C5700", bold=True), stopIfTrue=True))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($B{BV0}<>"",ISNUMBER({pct}))'],
    fill=fill("C6EFCE"), font=Font(color="006100", bold=True), stopIfTrue=True))
ws.conditional_formatting.add(f"H{BV0}:H{TR}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color=ACCENT_MID))

# --- Where the money goes: top 7 categories + everything else
WM_HEAD = TR + 3          # 35
put(ws, f"B{WM_HEAD}", "Where the money goes", size=13, bold=True, color=ACCENT)
WT = WM_HEAD + 1          # 36 table header
for col, hd in zip("JKLM", ["#", "Category", '="Spent ("&CurSym&")"', "Share"]):
    put(ws, f"{col}{WT}", hd, bold=True, color=WHITE, bg=ACCENT, h="center", border=BOX)
keys = f"$P${BV0}:$P${BV1}"
W0 = WT + 1
for k in range(1, 8):
    r = W0 + k - 1
    pos = f"MATCH(LARGE({keys},{k}),{keys},0)"
    put(ws, f"J{r}", k, h="center", color=MUTED, border=BOX)
    put(ws, f"K{r}", f'=IFERROR(INDEX($B${BV0}:$B${BV1},{pos}),"")', border=BOX)
    put(ws, f"L{r}", f'=IF(K{r}="","",INDEX($E${BV0}:$E${BV1},{pos}))', fmt=MONEY, border=BOX)
    put(ws, f"M{r}", f'=IF(K{r}="","",IFERROR(L{r}/$C$8,0))', fmt="0.0%", h="center", border=BOX)
OTH = W0 + 7
put(ws, f"J{OTH}", None, border=BOX)
put(ws, f"K{OTH}", "All other categories", italic=True, border=BOX)
put(ws, f"L{OTH}", f"=MAX(0,$C$8-SUM(L{W0}:L{OTH - 1}))", fmt=MONEY, border=BOX)
put(ws, f"M{OTH}", f"=IFERROR(L{OTH}/$C$8,0)", fmt="0.0%", h="center", border=BOX)
put(ws, f"J{OTH + 1}", None, bg=ACCENT_SOFT, border=BOX)
put(ws, f"K{OTH + 1}", "Total expenses", bold=True, bg=ACCENT_SOFT, border=BOX)
put(ws, f"L{OTH + 1}", f"=SUM(L{W0}:L{OTH})", fmt=MONEY, bold=True, bg=ACCENT_SOFT, border=BOX)
put(ws, f"M{OTH + 1}", f"=IFERROR(L{OTH + 1}/$C$8,0)", fmt="0.0%", h="center", bold=True,
    bg=ACCENT_SOFT, border=BOX)

# --- Day by day + Top 5
DAY_HEAD = WM_HEAD + 21   # 56
put(ws, f"B{DAY_HEAD}", "Day by day", size=13, bold=True, color=ACCENT)
TOP_HEAD = DAY_HEAD + 18  # 74
put(ws, f"B{TOP_HEAD}", "Top 5 biggest expenses this month", size=13, bold=True, color=ACCENT)
for col, hd in zip("BCDEFG", ["Description", "Date", "Category", '="Amount ("&CurSym&")"',
                              "Payment", "Share"]):
    put(ws, f"{col}{TOP_HEAD + 1}", hd, bold=True, color=WHITE, bg=ACCENT, h="center", border=BOX)
for k in range(1, 6):
    r = TOP_HEAD + 1 + k
    put(ws, f"P{r}", f'=IFERROR(MATCH(LARGE(tblTx[Rank Key],{k}),tblTx[Rank Key],0),"")', fmt="0")
    empty = '"No expenses logged yet"' if k == 1 else '""'
    put(ws, f"B{r}", f'=IF($P{r}="",{empty},IF(INDEX(tblTx[Description],$P{r})="","(no description)",'
                     f'INDEX(tblTx[Description],$P{r})))', border=BOX)
    put(ws, f"C{r}", f'=IF($P{r}="","",INDEX(tblTx[Date],$P{r}))', fmt="dd-mmm", h="center", border=BOX)
    put(ws, f"D{r}", f'=IF($P{r}="","",INDEX(tblTx[Category],$P{r})&"")', border=BOX)
    put(ws, f"E{r}", f'=IF($P{r}="","",INDEX(tblTx[Amount],$P{r}))', fmt=MONEY, bold=True, border=BOX)
    put(ws, f"F{r}", f'=IF($P{r}="","",INDEX(tblTx[Payment Method],$P{r})&"")', h="center", border=BOX)
    put(ws, f"G{r}", f'=IF($P{r}="","",IFERROR(E{r}/$C$8,0))', fmt="0.0%", h="center", border=BOX)

for hr in (10, WM_HEAD, DAY_HEAD, TOP_HEAD):
    ws.row_dimensions[hr].height = 22
ws.protection.sheet = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- charts


def style_axes(ch, y_fmt="#,##0"):
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.number_format = y_fmt
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E3E3E3")))
    for ax in (ch.x_axis, ch.y_axis):
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill="BFBFBF"))


def small_text(size):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=size)),
                                 endParaRPr=CharacterProperties(sz=size))])


def solid(series, color, line=False, width=28575):
    series.graphicalProperties = GraphicalProperties(solidFill=None if line else color)
    series.graphicalProperties.line.solidFill = color
    if line:
        series.graphicalProperties.line.width = width


# a) doughnut — where the money goes (top 7 + other)
dn = DoughnutChart(holeSize=55)
dn.title = "Where the money goes"
dn.add_data(Reference(ws, min_col=12, min_row=W0, max_row=OTH), titles_from_data=False)
dn.set_categories(Reference(ws, min_col=11, min_row=W0, max_row=OTH))
s = dn.series[0]
s.tx = SeriesLabel(v="Spent")
for idx, color in enumerate(DONUT_COLORS):
    pt = DataPoint(idx=idx)
    pt.graphicalProperties = GraphicalProperties(solidFill=color)
    pt.graphicalProperties.line.solidFill = WHITE
    pt.graphicalProperties.line.width = 19050
    s.dPt.append(pt)
dn.legend.position = "r"
dn.width, dn.height = 17.5, 8.8
ws.add_chart(dn, f"B{WT}")

# b) clustered bar — budget vs actual per category
bar = BarChart()
bar.type = "bar"
bar.grouping = "clustered"
bar.title = "Budget vs actual"
bar.add_data(Reference(ws, min_col=4, max_col=5, min_row=BV0 - 1, max_row=BV1), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=2, min_row=BV0, max_row=BV1))
solid(bar.series[0], "C5CCD6")
solid(bar.series[1], ACCENT_MID)
bar.gapWidth = 50
bar.overlap = -5
style_axes(bar)
bar.x_axis.scaling.orientation = "maxMin"   # first category at the top
bar.y_axis.crosses = "max"                  # keep the value axis at the bottom
bar.x_axis.tickLblSkip = 1
bar.x_axis.txPr = small_text(800)
bar.legend.position = "b"
bar.width, bar.height = 14.2, 11.3
ws.add_chart(bar, "J10")

# c) line — running balance day by day (date axis, fed by Daily Flow)
fl = ws_flow
line = LineChart()
line.title = "Running balance through the month"
line.add_data(Reference(fl, min_col=9, min_row=F0 - 1, max_row=F1), titles_from_data=True)
line.set_categories(Reference(fl, min_col=8, min_row=F0, max_row=F1))
ls = line.series[0]
solid(ls, ACCENT_MID, line=True)
ls.marker = Marker(symbol="none")
ls.smooth = False
line.x_axis = DateAxis(crossAx=100)
line.x_axis.number_format = "d"
line.x_axis.majorTimeUnit = "days"
line.x_axis.baseTimeUnit = "days"
line.y_axis.crossAx = 500
style_axes(line)
line.legend = None
line.width, line.height = 19.8, 8
ws.add_chart(line, f"B{DAY_HEAD + 1}")

# d) column — daily spending
col = BarChart()
col.type = "col"
col.title = "Daily spending"
col.add_data(Reference(fl, min_col=10, min_row=F0 - 1, max_row=F1), titles_from_data=True)
col.set_categories(Reference(fl, min_col=8, min_row=F0, max_row=F1))
solid(col.series[0], ACCENT_MID)
col.gapWidth = 40
col.x_axis = DateAxis(crossAx=100)
col.x_axis.number_format = "d"
col.x_axis.majorTimeUnit = "days"
col.x_axis.baseTimeUnit = "days"
col.y_axis.crossAx = 500
style_axes(col)
col.legend = None
col.width, col.height = 14.2, 8
ws.add_chart(col, f"J{DAY_HEAD + 1}")

# ================================================================ START HERE
ws = ws_start
ws.sheet_view.showGridLines = False
widths(ws, {"A": 2, "B": 6, "C": 100})
put(ws, "B2", "Monthly Budget Tracker", size=22, bold=True, color=ACCENT)
put(ws, "B3", "See where your money goes, how much is left, and whether you're on budget — "
              "from the 1st to the last day of the month.", color=MUTED, italic=True)
put(ws, "B5", "How to use it", size=14, bold=True, color=ACCENT)
steps = [
    ("Set the month",
     "Open the Setup sheet. In the yellow 'Month start date' cell type the 1st day of the month "
     "(e.g. 1-Nov-2026) and choose your currency symbol. Dates, day counts and charts adjust by themselves."),
    ("Enter your income and budgets",
     "Still on Setup: list your income sources with the amount you expect, then give each expense "
     "category a monthly budget. Rename, add or remove rows freely — the dropdowns follow."),
    ("Log your expenses every day",
     "On Transactions, add one row per income or expense: Date, Type, Category, Description, Amount, "
     "Payment method. Pick Type first — the Category list then shows income sources or expense categories. "
     "First delete the 20 sample rows: select A5:G24 and press Delete."),
    ("Check the Dashboard",
     "Totals, what's left, % of income spent, budget vs actual for every category, charts and your "
     "5 biggest expenses — all update on their own. Daily Flow shows the numbers behind the daily charts."),
]
r = 6
for n, (head, body) in enumerate(steps, 1):
    put(ws, f"B{r}", n, size=14, bold=True, color=WHITE, bg=ACCENT, h="center")
    put(ws, f"C{r}", head, size=11, bold=True)
    put(ws, f"C{r + 1}", body, wrap=True, v="top", color=INK)
    ws.row_dimensions[r].height = 22
    ws.row_dimensions[r + 1].height = 15 * (len(body) // 95 + 1) + 3
    ws.merge_cells(f"B{r}:B{r + 1}")
    r += 3

put(ws, f"B{r}", "Color legend", size=14, bold=True, color=ACCENT)
legend = [
    (INPUT_FILL, "Light yellow = input cells. Type or change these."),
    (AUTO_FILL, "Grey = automatic. Don't type over these — they contain formulas."),
    (ACCENT, "Blue = headings and table headers."),
]
for color, text in legend:
    r += 1
    put(ws, f"B{r}", None, bg=color, border=BOX)
    put(ws, f"C{r}", text)
    ws.row_dimensions[r].height = 20

r += 2
put(ws, f"B{r}", "Good to know", size=14, bold=True, color=ACCENT)
tips = [
    "Add an income source, category or payment method by typing in the empty row directly under its table on Setup.",
    "The Dashboard shows up to 20 expense categories (15 are pre-filled, so 5 slots are free).",
    "Transactions dated outside the selected month are rejected when typed, and ignored by the Dashboard if pasted in.",
    "New month? Save a copy of the file, clear the yellow cells on Transactions, and change the Month start date on Setup.",
    "Dashboard and Daily Flow are protected against accidental edits (no password: Review ▸ Unprotect Sheet).",
    "'Days left' and 'Average daily spend' use today's date while the month is in progress, and the full month otherwise.",
]
for t in tips:
    r += 1
    put(ws, f"B{r}", "•", h="center", color=ACCENT, bold=True)
    put(ws, f"C{r}", t, wrap=True)
    ws.row_dimensions[r].height = 15 * (len(t) // 95 + 1) + 3

wb.active = 0
wb.save(OUT)
print("saved", OUT)


# ================================================================ cache computed values
def cache_values(path):
    """Recalculate a copy with LibreOffice and write the results into `path` as cached values."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found - skipping cached values (Excel will still calculate on open)")
        return
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "in.xlsx")
        shutil.copy(path, src)
        outdir = os.path.join(tmp, "out")
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", outdir, src],
                       check=True, capture_output=True, timeout=180)
        calc_path = os.path.join(outdir, "in.xlsx")

        def sheet_files(z):
            wbx = etree.fromstring(z.read("xl/workbook.xml"))
            rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            target = {r.get("Id"): r.get("Target") for r in rels}
            out = {}
            for s in wbx.find("m:sheets", ns):
                t = target[s.get(f"{{{ns['r']}}}id")].lstrip("/")
                out[s.get("name")] = t if t.startswith("xl/") else "xl/" + t
            return out

        with zipfile.ZipFile(calc_path) as zc:
            shared = []
            if "xl/sharedStrings.xml" in zc.namelist():
                for si in etree.fromstring(zc.read("xl/sharedStrings.xml")).findall("m:si", ns):
                    shared.append("".join(si.itertext()))
            values = {}
            for name, f in sheet_files(zc).items():
                for c in etree.fromstring(zc.read(f)).iter(f"{{{ns['m']}}}c"):
                    if c.find("m:f", ns) is None:
                        continue
                    v = c.find("m:v", ns)
                    t = c.get("t", "n")
                    val = "" if v is None or v.text is None else v.text
                    if t == "s":
                        t, val = "str", shared[int(val)]
                    values[(name, c.get("r"))] = (t, val)

        tmp_out = os.path.join(tmp, "final.xlsx")
        with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp_out, "w", zipfile.ZIP_DEFLATED) as zo:
            files = {v: k for k, v in sheet_files(zi).items()}
            for item in zi.infolist():
                data = zi.read(item.filename)
                if item.filename in files:
                    sheet = files[item.filename]
                    root = etree.fromstring(data)
                    for c in root.iter(f"{{{ns['m']}}}c"):
                        if c.find("m:f", ns) is None or (sheet, c.get("r")) not in values:
                            continue
                        t, val = values[(sheet, c.get("r"))]
                        v = c.find("m:v", ns)
                        if v is None:
                            v = etree.SubElement(c, f"{{{ns['m']}}}v")
                        v.text = val
                        if t == "n":
                            c.attrib.pop("t", None)
                        else:
                            c.set("t", t)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                zo.writestr(item, data)
        shutil.move(tmp_out, path)
    print("cached values written from LibreOffice recalculation")


cache_values(OUT)
